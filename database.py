"""
SQLite storage for the Pricing App.

Two ways to keep prices current (both supported):
  * Re-import: reload everything from the pricing workbook (openpyxl).
  * Admin edit: update an individual product price / adder in the DB.

Stdlib only (sqlite3). openpyxl is used solely for the workbook import.
"""
import glob
import json
import os
import sqlite3
from datetime import datetime, timezone

import pricing_engine as pe

HERE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.path.join(HERE, "pricing.db")


def latest_pricing_workbook():
    """Newest 'Pricing *.xlsx' in the app folder (ignores Excel ~$ lock files).

    Lets a user drop in a re-dated file (Pricing 7.14.26.xlsx -> Pricing 9.9.26.xlsx)
    without editing any code.
    """
    files = [p for p in glob.glob(os.path.join(HERE, "Pricing *.xlsx"))
             if not os.path.basename(p).startswith("~$")]
    if not files:
        return os.path.join(HERE, "Pricing.xlsx")  # nonexistent -> clear error on import
    return max(files, key=os.path.getmtime)


DEFAULT_WORKBOOK = latest_pricing_workbook()

# Outbound freight by destination. cwt assumes ~full truckload (all_in / 46,000 lb x 100).
# Regions each carry a FREIGHT $/cwt and a COMPETITIVE $/cwt. On a quote the rep applies
# freight only, competitive only, or both (they stack). freight_cwt = pass-through cost
# (excluded from CRU Spread); comp_cwt = a price move (counts toward margin).
# NOTE: DFW / East TX / Houston-South TX freight and ALL comp_cwt are PLACEHOLDERS - confirm.
REGIONS = [
    {"region": "DFW",              "freight_cwt": 0.75, "comp_cwt": 0.00},
    {"region": "East TX",          "freight_cwt": 1.25, "comp_cwt": 0.00},
    {"region": "OKC",              "freight_cwt": 1.70, "comp_cwt": 0.00},
    {"region": "Tulsa",            "freight_cwt": 2.05, "comp_cwt": 0.00},
    {"region": "Houston/South TX", "freight_cwt": 2.00, "comp_cwt": 0.00},
    {"region": "West TX",          "freight_cwt": 2.55, "comp_cwt": 0.00},
]

SCHEMA = """
CREATE TABLE IF NOT EXISTS products (
    key         TEXT PRIMARY KEY,
    thickness   TEXT NOT NULL,
    thickness_in REAL,
    grade       TEXT NOT NULL,
    width       TEXT,
    width_num   REAL,
    price_cwt   REAL,
    price_note  TEXT
);
CREATE TABLE IF NOT EXISTS stock_tiers  (tier_label TEXT, max_weight INTEGER, adder_cwt REAL);
CREATE TABLE IF NOT EXISTS length_tiers (tier_label TEXT, max_weight INTEGER, adder_cwt REAL);
CREATE TABLE IF NOT EXISTS extras (name TEXT PRIMARY KEY, adder_cwt REAL, selectable INTEGER DEFAULT 0);
CREATE TABLE IF NOT EXISTS meta (k TEXT PRIMARY KEY, v TEXT);
CREATE TABLE IF NOT EXISTS freight (destination TEXT PRIMARY KEY, freight_cwt REAL, comp_cwt REAL);
CREATE TABLE IF NOT EXISTS quotes (
    id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT, customer TEXT, quote_no TEXT,
    destination TEXT, adder_basis TEXT, use_inventory INTEGER, grand_total REAL, payload TEXT
);
CREATE TABLE IF NOT EXISTS customers (
    name TEXT PRIMARY KEY, segment TEXT, playbook INTEGER, credit TEXT,
    dyn_cwt REAL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS cust_config (k TEXT PRIMARY KEY, v REAL);
"""

# Customer pricing layers ($/cwt), tunable in Admin. Playbook 1/2/3, credit High/Low,
# and a DYNAMIC margin that ratchets up per win (capped) and eases back per loss.
CUSTOMER_CONFIG = {"pb1_cwt": -2.0, "pb2_cwt": 0.0, "pb3_cwt": 2.0,
                   "credit_high_cwt": 2.0, "credit_low_cwt": 0.0,
                   "dyn_step_cwt": 0.5, "dyn_cap_cwt": 3.0, "dyn_loss_cwt": 0.5}


def connect():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = connect()
    conn.executescript(SCHEMA)
    conn.commit()
    conn.close()


def _s(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _tier_weight(label):
    digits = "".join(ch for ch in str(label) if ch.isdigit())
    return int(digits) if digits else None


def import_from_workbook(path=None):
    """Wipe and reload all pricing data from the Excel workbook.

    With no path, always picks the newest 'Pricing *.xlsx' currently in the folder.
    """
    import openpyxl

    if path is None:
        path = latest_pricing_workbook()
    if not os.path.exists(path):
        raise FileNotFoundError(path)

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["1011-36"]

    products, seen = [], set()
    for r in range(2, ws.max_row + 1):
        th, gr, wd, price = (ws.cell(r, 1).value, ws.cell(r, 2).value,
                             ws.cell(r, 3).value, ws.cell(r, 5).value)
        if th is None and gr is None:
            continue
        th_s, gr_s, wd_s = _s(th), _s(gr), _s(wd)
        key = "-".join(part for part in (th_s, wd_s, gr_s) if part)
        if key in seen:      # keep first occurrence
            continue
        seen.add(key)
        numeric = isinstance(price, (int, float))
        products.append((key, th_s, pe.thickness_to_inch(th_s), gr_s, wd_s,
                         pe.width_to_number(wd_s),
                         float(price) if numeric else None,
                         None if numeric else _s(price)))

    stock, length, target = [], [], None
    for r in range(3, 20):
        g, h = ws.cell(r, 7).value, ws.cell(r, 8).value
        if g is None:
            continue
        gl = str(g).lower()
        if "stock adders" in gl:
            target = stock; continue
        if "custom length" in gl:
            target = length; continue
        if target is not None and isinstance(h, (int, float)):
            target.append((_s(g), _tier_weight(g), float(h)))

    cost = wb["Cost"]
    extras = []
    for r in range(2, cost.max_row + 1):
        d, e = cost.cell(r, 4).value, cost.cell(r, 5).value
        if d is not None and isinstance(e, (int, float)):
            name = _s(d)
            selectable = 1 if name in pe.ALLOWED_EXTRAS else 0
            extras.append((name, float(e), selectable))

    conn = connect()
    conn.executescript(SCHEMA)
    for t in ("products", "stock_tiers", "length_tiers", "extras"):
        conn.execute(f"DELETE FROM {t}")
    conn.executemany("INSERT INTO products VALUES (?,?,?,?,?,?,?,?)", products)
    conn.executemany("INSERT INTO stock_tiers VALUES (?,?,?)", stock)
    conn.executemany("INSERT INTO length_tiers VALUES (?,?,?)", length)
    conn.executemany("INSERT OR REPLACE INTO extras VALUES (?,?,?)", extras)
    conn.execute("INSERT OR REPLACE INTO meta VALUES ('last_import', ?)",
                 (datetime.now(timezone.utc).isoformat(timespec="seconds"),))
    conn.execute("INSERT OR REPLACE INTO meta VALUES ('source_workbook', ?)", (os.path.basename(path),))
    conn.commit()
    conn.close()
    return {"products": len(products), "stock_tiers": len(stock),
            "length_tiers": len(length), "extras": len(extras)}


# ---- freight ------------------------------------------------------------
def round_005(x):
    """Round a $/cwt value to the nearest $0.05."""
    return round(round(float(x) / 0.05) * 0.05, 2)


def ensure_freight():
    conn = connect()
    conn.executescript(SCHEMA)
    reseed = False
    try:
        current = {r[0] for r in conn.execute("SELECT destination FROM freight WHERE freight_cwt IS NOT NULL")}
        if current != {r["region"] for r in REGIONS}:
            reseed = True   # regions added/removed
    except sqlite3.OperationalError:
        reseed = True       # old schema (no freight_cwt column)
    if reseed:
        conn.execute("DROP TABLE IF EXISTS freight")
        conn.executescript(SCHEMA)
        conn.executemany("INSERT OR REPLACE INTO freight VALUES (?,?,?)",
                         [(r["region"], round_005(r["freight_cwt"]), round_005(r["comp_cwt"])) for r in REGIONS])
        conn.commit()
    conn.close()


def get_freight(destination):
    ensure_freight()
    conn = connect()
    row = conn.execute("SELECT * FROM freight WHERE destination=?", (destination,)).fetchone()
    conn.close()
    return dict(row) if row else None


# ---- read helpers -------------------------------------------------------
def bootstrap():
    ensure_freight()
    """Everything the UI needs to render the form."""
    conn = connect()
    prods = [dict(r) for r in conn.execute(
        "SELECT key, thickness, thickness_in, grade, width, width_num, price_cwt, price_note FROM products "
        "ORDER BY thickness, grade, width")]
    extras = [dict(r) for r in conn.execute(
        "SELECT name, adder_cwt FROM extras WHERE selectable=1 ORDER BY name")]
    meta = {r["k"]: r["v"] for r in conn.execute("SELECT k, v FROM meta")}
    freight = [dict(r) for r in conn.execute("SELECT destination, freight_cwt, comp_cwt FROM freight ORDER BY destination")]
    conn.close()
    return {"products": prods, "extras": extras, "meta": meta, "freight": freight}


def get_product(key):
    conn = connect()
    row = conn.execute("SELECT * FROM products WHERE key=?", (key,)).fetchone()
    conn.close()
    return dict(row) if row else None


def get_tiers():
    conn = connect()
    stock = [dict(r) for r in conn.execute("SELECT * FROM stock_tiers")]
    length = [dict(r) for r in conn.execute("SELECT * FROM length_tiers")]
    extras = {r["name"]: r["adder_cwt"] for r in conn.execute("SELECT name, adder_cwt FROM extras")}
    conn.close()
    return stock, length, extras


def update_price(key, price_cwt):
    conn = connect()
    cur = conn.execute("UPDATE products SET price_cwt=?, price_note=NULL WHERE key=?",
                       (float(price_cwt), key))
    conn.commit()
    changed = cur.rowcount
    conn.close()
    return changed


def save_quote(customer, quote_no, quote, grand_total):
    conn = connect()
    conn.executescript(SCHEMA)
    cur = conn.execute(
        "INSERT INTO quotes (created_at, customer, quote_no, destination, adder_basis, use_inventory, grand_total, payload) "
        "VALUES (?,?,?,?,?,?,?,?)",
        (datetime.now().isoformat(timespec="seconds"), (customer or "").strip(), (quote_no or "").strip(),
         quote.get("destination", ""), quote.get("adder_basis", "line"),
         1 if quote.get("use_inventory") else 0, float(grand_total or 0), json.dumps(quote)))
    conn.commit()
    qid = cur.lastrowid
    conn.close()
    return qid


def list_quotes(limit=50):
    _ensure_quote_status()
    conn = connect()
    conn.executescript(SCHEMA)
    out = []
    for r in conn.execute("SELECT id, created_at, customer, quote_no, destination, grand_total, "
                          "COALESCE(status,'open') status, payload "
                          "FROM quotes ORDER BY id DESC LIMIT ?", (limit,)):
        d = dict(r)
        try:
            d["lines"] = len(json.loads(d.pop("payload")).get("lines", []))
        except Exception:
            d.pop("payload", None); d["lines"] = 0
        out.append(d)
    conn.close()
    return out


def get_quote(qid):
    conn = connect()
    conn.executescript(SCHEMA)
    row = conn.execute("SELECT * FROM quotes WHERE id=?", (qid,)).fetchone()
    conn.close()
    if not row:
        return None
    d = dict(row)
    try:
        d["payload"] = json.loads(d["payload"])
    except Exception:
        d["payload"] = {}
    return d


def delete_quote(qid):
    conn = connect()
    conn.executescript(SCHEMA)
    n = conn.execute("DELETE FROM quotes WHERE id=?", (qid,)).rowcount
    conn.commit()
    conn.close()
    return n


# ---- customers ----------------------------------------------------------
def get_cust_config():
    conn = connect()
    conn.executescript(SCHEMA)
    have = {r["k"]: r["v"] for r in conn.execute("SELECT k, v FROM cust_config")}
    for k, v in CUSTOMER_CONFIG.items():
        if k not in have:
            conn.execute("INSERT INTO cust_config VALUES (?,?)", (k, v))
    conn.commit()
    cfg = {r["k"]: r["v"] for r in conn.execute("SELECT k, v FROM cust_config")}
    conn.close()
    return {**CUSTOMER_CONFIG, **cfg}


def set_cust_config(updates):
    conn = connect()
    conn.executescript(SCHEMA)
    for k in CUSTOMER_CONFIG:
        if k in updates and updates[k] is not None:
            conn.execute("INSERT OR REPLACE INTO cust_config VALUES (?,?)", (k, float(updates[k])))
    conn.commit()
    conn.close()
    return get_cust_config()


def list_customers():
    _ensure_quote_status()   # guarantees the dyn_cwt column on pre-existing DBs
    conn = connect()
    conn.executescript(SCHEMA)
    rows = [dict(r) for r in conn.execute(
        "SELECT name, segment, playbook, credit, COALESCE(dyn_cwt,0) dyn_cwt "
        "FROM customers ORDER BY name")]
    conn.close()
    return rows


def get_customer(name):
    if not name:
        return None
    conn = connect()
    conn.executescript(SCHEMA)
    row = conn.execute("SELECT * FROM customers WHERE name=?", (name.strip(),)).fetchone()
    conn.close()
    return dict(row) if row else None


def upsert_customer(name, segment=None, playbook=2, credit="low"):
    name = (name or "").strip()
    if not name:
        raise ValueError("customer name required")
    pb = int(playbook) if playbook in (1, 2, 3, "1", "2", "3") else 2
    credit = "high" if str(credit).lower().startswith("h") else "low"
    _ensure_quote_status()   # dyn_cwt column
    conn = connect()
    conn.executescript(SCHEMA)
    # Upsert without clobbering the dynamic margin the win/loss history has built up.
    conn.execute(
        "INSERT INTO customers (name, segment, playbook, credit) VALUES (?,?,?,?) "
        "ON CONFLICT(name) DO UPDATE SET segment=excluded.segment, "
        "playbook=excluded.playbook, credit=excluded.credit",
        (name, (segment or "").strip(), pb, credit))
    conn.commit()
    conn.close()
    return get_customer(name)


def delete_customer(name):
    conn = connect()
    conn.executescript(SCHEMA)
    n = conn.execute("DELETE FROM customers WHERE name=?", ((name or "").strip(),)).rowcount
    conn.commit()
    conn.close()
    return n


def customer_adjustments(name):
    """List of {label, cwt} $/cwt layers for a customer's playbook + credit (stacks)."""
    c = get_customer(name)
    if not c:
        return []
    cfg = get_cust_config()
    adjs = []
    pb_amt = cfg.get("pb%d_cwt" % (c["playbook"] or 2), 0.0)
    if pb_amt:
        adjs.append({"label": "Playbook %d" % c["playbook"], "cwt": pb_amt})
    if (c["credit"] or "low") == "high" and cfg.get("credit_high_cwt"):
        adjs.append({"label": "Credit: High Risk", "cwt": cfg["credit_high_cwt"]})
    dyn = c.get("dyn_cwt") or 0.0
    if dyn:
        adjs.append({"label": "Dynamic margin (won history)", "cwt": dyn})
    return adjs


CUST_DB_CONN = os.environ.get(
    "PRICING_CUST_DB_CONN",
    "DRIVER={SQL Server};SERVER=10.0.1.50;DATABASE=Planning;Trusted_Connection=yes")
BOOKINGS_DB_CONN = os.environ.get(
    "PRICING_BOOKINGS_DB_CONN",
    "DRIVER={SQL Server};SERVER=10.0.1.50;DATABASE=Inventory_Planning;Trusted_Connection=yes")
DASH_DEFAULT = {"conv_alert_pct": 60.0}


def _ensure_quote_status():
    conn = connect()
    conn.executescript(SCHEMA)
    qcols = {r[1] for r in conn.execute("PRAGMA table_info(quotes)")}
    if "status" not in qcols:
        conn.execute("ALTER TABLE quotes ADD COLUMN status TEXT DEFAULT 'open'")
    if "lost_reason" not in qcols:
        conn.execute("ALTER TABLE quotes ADD COLUMN lost_reason TEXT")
    if "competitor_cwt" not in qcols:
        conn.execute("ALTER TABLE quotes ADD COLUMN competitor_cwt REAL")
    ccols = {r[1] for r in conn.execute("PRAGMA table_info(customers)")}
    if "dyn_cwt" not in ccols:
        conn.execute("ALTER TABLE customers ADD COLUMN dyn_cwt REAL DEFAULT 0")
    conn.execute("CREATE TABLE IF NOT EXISTS dash_config (k TEXT PRIMARY KEY, v REAL)")
    conn.commit()
    conn.close()


def _apply_dynamic(name, outcome):
    """Ratchet a customer's dynamic margin: +step per win (capped), -backoff per loss."""
    cfg = get_cust_config()
    c = get_customer(name)
    if not c:
        return
    dyn = c.get("dyn_cwt") or 0.0
    if outcome == "won":
        dyn = min(dyn + cfg["dyn_step_cwt"], cfg["dyn_cap_cwt"])
    elif outcome == "lost":
        dyn = max(dyn - cfg["dyn_loss_cwt"], 0.0)
    conn = connect()
    conn.executescript(SCHEMA)
    conn.execute("UPDATE customers SET dyn_cwt=? WHERE name=?", (round(dyn, 2), name))
    conn.commit()
    conn.close()


LOSS_REASONS = ("price", "lead_time", "availability", "relationship", "no_bid", "other")


def set_quote_status(qid, status, lost_reason=None, competitor_cwt=None):
    _ensure_quote_status()
    status = status if status in ("open", "won", "lost") else "open"
    reason = lost_reason if (status == "lost" and lost_reason in LOSS_REASONS) else None
    try:
        comp = float(competitor_cwt) if (status == "lost" and competitor_cwt not in (None, "")) else None
    except (TypeError, ValueError):
        comp = None
    conn = connect()
    row = conn.execute("SELECT COALESCE(status,'open'), payload FROM quotes WHERE id=?", (qid,)).fetchone()
    if not row:
        conn.close(); return 0
    prior = row[0]
    # Clear stale loss feedback when a quote moves off 'lost'.
    conn.execute("UPDATE quotes SET status=?, lost_reason=?, competitor_cwt=? WHERE id=?",
                 (status, reason, comp, qid))
    conn.commit()
    conn.close()
    if status != prior and status in ("won", "lost"):   # only on a real transition
        try:
            name = json.loads(row[1]).get("customer_name")
        except Exception:
            name = None
        if name:
            _apply_dynamic(name, status)
    return 1


def get_dash_config():
    _ensure_quote_status()
    conn = connect()
    cfg = {r["k"]: r["v"] for r in conn.execute("SELECT k, v FROM dash_config")}
    conn.close()
    return {**DASH_DEFAULT, **cfg}


def set_dash_config(updates):
    _ensure_quote_status()
    conn = connect()
    for k in DASH_DEFAULT:
        if k in updates and updates[k] is not None:
            conn.execute("INSERT OR REPLACE INTO dash_config VALUES (?,?)", (k, float(updates[k])))
    conn.commit()
    conn.close()
    return get_dash_config()


def _quote_conversion():
    """Conversion from the app's own saved quotes (needs reps marking won/lost)."""
    _ensure_quote_status()
    conn = connect()
    rows = [dict(r) for r in conn.execute(
        "SELECT COALESCE(status,'open') status, COALESCE(destination,'(none)') region, COUNT(*) n, SUM(grand_total) val "
        "FROM quotes GROUP BY COALESCE(status,'open'), COALESCE(destination,'(none)')")]
    conn.close()
    tot = {"won": 0, "lost": 0, "open": 0}
    by_region = {}
    for r in rows:
        tot[r["status"]] = tot.get(r["status"], 0) + r["n"]
        reg = by_region.setdefault(r["region"], {"won": 0, "lost": 0, "open": 0})
        reg[r["status"]] += r["n"]
    def rate(w, l):
        return round(100.0 * w / (w + l), 1) if (w + l) else None
    regions = [{"region": k, "won": v["won"], "lost": v["lost"], "open": v["open"], "rate": rate(v["won"], v["lost"])}
               for k, v in sorted(by_region.items())]
    return {"won": tot["won"], "lost": tot["lost"], "open": tot["open"],
            "rate": rate(tot["won"], tot["lost"]), "by_region": regions}


def _loss_feedback():
    """Why the app's saved quotes were lost, and competitor $/cwt reps cited (beta learning)."""
    _ensure_quote_status()
    conn = connect()
    reasons = [dict(r) for r in conn.execute(
        "SELECT COALESCE(lost_reason,'(unspecified)') reason, COUNT(*) n "
        "FROM quotes WHERE status='lost' GROUP BY COALESCE(lost_reason,'(unspecified)') "
        "ORDER BY n DESC")]
    comp = [dict(r) for r in conn.execute(
        "SELECT COALESCE(destination,'(none)') region, COUNT(*) n, "
        "ROUND(AVG(competitor_cwt),2) avg_comp, ROUND(MIN(competitor_cwt),2) min_comp, "
        "ROUND(MAX(competitor_cwt),2) max_comp "
        "FROM quotes WHERE status='lost' AND competitor_cwt IS NOT NULL "
        "GROUP BY COALESCE(destination,'(none)') ORDER BY n DESC")]
    conn.close()
    return {"reasons": reasons, "competitor_by_region": comp}


def _norm_name(s):
    return " ".join((s or "").upper().split())


_NAME_SUFFIX = set("INC LLC CO CORP CORPORATION LP LTD LLP USA THE COMPANY "
                   "INDUSTRIES IND MFG MANUFACTURING".split())


def _aggr_tokens(s):
    """Aggressively-normalized significant tokens: drop punctuation + corporate suffixes."""
    import re
    s = re.sub(r"[^A-Z0-9 ]", " ", (s or "").upper())
    return [t for t in s.split() if t and t not in _NAME_SUFFIX]


def erp_name_map():
    """{cus_id -> customer name} from the Salesforce-fed table (matches sahstn_rec.stn_sld_cus_id)."""
    import pyodbc
    cn = pyodbc.connect(CUST_DB_CONN, timeout=20); cur = cn.cursor()
    cur.execute("SELECT LTRIM(RTRIM(bka_sld_cus_id)) cid, MAX(LTRIM(RTRIM(customer))) nm "
                "FROM dbo.salesforce_sales_numbers WHERE customer IS NOT NULL "
                "GROUP BY LTRIM(RTRIM(bka_sld_cus_id))")
    m = {r[0]: r[1] for r in cur.fetchall()}
    cn.close()
    return m


def erp_segment_map():
    """Dominant dim_seg per customer from salesforce_billings.
    Returns (by_id, resolve) where by_id maps cus_id -> segment and
    resolve(name) -> segment using exact -> suffix-stripped -> safe token-subset
    name matching (token-subset only when it points to a single segment)."""
    import pyodbc
    cn = pyodbc.connect(CUST_DB_CONN, timeout=30); cur = cn.cursor()
    cur.execute("""WITH x AS (
        SELECT LTRIM(RTRIM(cust_id)) cid, LTRIM(RTRIM(customer)) nm, LTRIM(RTRIM(dim_seg)) seg,
               COUNT(*) n,
               ROW_NUMBER() OVER (PARTITION BY LTRIM(RTRIM(cust_id))
                                  ORDER BY COUNT(*) DESC) rn
        FROM dbo.salesforce_billings
        WHERE dim_seg IS NOT NULL AND LTRIM(RTRIM(dim_seg)) <> '' AND cust_id IS NOT NULL
        GROUP BY LTRIM(RTRIM(cust_id)), LTRIM(RTRIM(customer)), LTRIM(RTRIM(dim_seg)))
        SELECT cid, nm, seg FROM x WHERE rn = 1""")
    by_id, exact, aggr = {}, {}, {}
    tok_index = []   # (set(tokens), seg)
    for cid, nm, seg in cur.fetchall():
        by_id[cid] = seg
        if not nm:
            continue
        exact.setdefault(_norm_name(nm), seg)
        toks = _aggr_tokens(nm)
        aggr.setdefault(" ".join(toks), seg)
        if toks:
            tok_index.append((set(toks), seg))
    cn.close()

    def resolve(name):
        e = exact.get(_norm_name(name))
        if e:
            return e
        toks = _aggr_tokens(name)
        if not toks:
            return None
        a = aggr.get(" ".join(toks))
        if a:
            return a
        mine = set(toks)
        segs = {seg for tset, seg in tok_index if mine <= tset}   # our name fully inside theirs
        return next(iter(segs)) if len(segs) == 1 else None

    return by_id, resolve


def _bookings(days=90):
    """Bookings from the ERP sales history (sahstn_rec); names via salesforce_sales_numbers,
    segment via salesforce_billings.dim_seg."""
    import pyodbc
    cn = pyodbc.connect(BOOKINGS_DB_CONN, timeout=30)
    cur = cn.cursor()
    cur.execute("SELECT SUM(stn_tot_val), SUM(stn_blg_wgt), COUNT(*) FROM dbo.sahstn_rec "
                "WHERE stn_shp_dt >= DATEADD(day, ?, GETDATE())", -abs(days))
    val, wgt, cnt = cur.fetchone()
    cur.execute("SELECT FORMAT(stn_shp_dt,'yyyy-MM') ym, SUM(stn_tot_val), SUM(stn_blg_wgt) "
                "FROM dbo.sahstn_rec WHERE stn_shp_dt >= DATEADD(month,-6,GETDATE()) "
                "GROUP BY FORMAT(stn_shp_dt,'yyyy-MM') ORDER BY ym")
    trend = [{"month": r[0], "val": float(r[1] or 0), "wgt": float(r[2] or 0)} for r in cur.fetchall()]
    cur.execute("SELECT TOP 10 LTRIM(RTRIM(stn_sld_cus_id)) cid, SUM(stn_tot_val) val, SUM(stn_blg_wgt) wgt "
                "FROM dbo.sahstn_rec WHERE stn_shp_dt >= DATEADD(day, ?, GETDATE()) "
                "GROUP BY LTRIM(RTRIM(stn_sld_cus_id)) ORDER BY SUM(stn_tot_val) DESC", -abs(days))
    top = [{"cid": r[0], "val": float(r[1] or 0), "wgt": float(r[2] or 0)} for r in cur.fetchall()]
    cn.close()
    # names via salesforce_sales_numbers, segment via salesforce_billings.dim_seg
    names, seg_by_id = {}, {}
    try:
        names = erp_name_map()
    except Exception:
        pass
    try:
        seg_by_id, _ = erp_segment_map()
    except Exception:
        pass
    for t in top:
        t["name"] = names.get(t["cid"], t["cid"])
        t["segment"] = seg_by_id.get(t["cid"], "")
    return {"days": days, "total_val": float(val or 0), "total_wgt": float(wgt or 0),
            "shipments": int(cnt or 0), "trend": trend, "top_customers": top}


def erp_quote_history():
    """Historical quotes from the ERP order file (ORTORH, prefix 'QT').
    Volume by salesperson and customer, plus a coarse customer-level conversion
    (share of quoted customers that also placed an order). NOTE: the extract has
    no quoted price or reliable date, so this is volume/coverage, not a win-rate."""
    import pyodbc
    cn = pyodbc.connect(CUST_DB_CONN, timeout=40); cur = cn.cursor()
    cur.execute("SELECT LTRIM(RTRIM(orh_ord_pfx)) pfx, COUNT(*) n FROM dbo.ORTORH "
                "GROUP BY LTRIM(RTRIM(orh_ord_pfx))")
    totals = {r[0]: int(r[1]) for r in cur.fetchall()}
    cur.execute("SELECT TOP 12 LTRIM(RTRIM(orh_tkn_slp)) slp, COUNT(*) n FROM dbo.ORTORH "
                "WHERE orh_ord_pfx='QT' GROUP BY LTRIM(RTRIM(orh_tkn_slp)) ORDER BY n DESC")
    by_rep = [{"rep": r[0] or "(none)", "quotes": int(r[1])} for r in cur.fetchall()]
    cur.execute("SELECT TOP 12 LTRIM(RTRIM(orh_sld_cus_id)) cid, COUNT(*) n FROM dbo.ORTORH "
                "WHERE orh_ord_pfx='QT' GROUP BY LTRIM(RTRIM(orh_sld_cus_id)) ORDER BY n DESC")
    top_rows = [(r[0], int(r[1])) for r in cur.fetchall()]
    cur.execute("""SELECT COUNT(DISTINCT q.orh_sld_cus_id), COUNT(DISTINCT s.cid)
        FROM (SELECT DISTINCT orh_sld_cus_id FROM dbo.ORTORH WHERE orh_ord_pfx='QT') q
        LEFT JOIN (SELECT DISTINCT orh_sld_cus_id cid FROM dbo.ORTORH WHERE orh_ord_pfx='SO') s
          ON s.cid=q.orh_sld_cus_id""")
    quoted, ordered = cur.fetchone()
    cn.close()
    names = {}
    try:
        names = erp_name_map()
    except Exception:
        pass
    by_customer = [{"name": names.get(cid, cid), "quotes": n} for cid, n in top_rows]
    quoted, ordered = int(quoted or 0), int(ordered or 0)
    return {"quotes": totals.get("QT", 0), "orders": totals.get("SO", 0),
            "by_rep": by_rep, "by_customer": by_customer,
            "customers_quoted": quoted, "customers_ordered": ordered,
            "coverage_pct": round(100.0 * ordered / quoted, 0) if quoted else None}


def salesperson_names():
    """{slp code -> 'First Last'} derived from salesperson_emails (email local part)."""
    import pyodbc
    out = {}
    try:
        cn = pyodbc.connect(CUST_DB_CONN, timeout=15); cur = cn.cursor()
        cur.execute("SELECT LTRIM(RTRIM(slp_slp)), usr_email FROM dbo.salesperson_emails "
                    "WHERE slp_slp IS NOT NULL")
        for code, email in cur.fetchall():
            local = (email or "").split("@")[0]
            parts = [p for p in local.replace("_", ".").split(".") if p]
            out[code] = " ".join(p.capitalize() for p in parts) if parts else code
        cn.close()
    except Exception:
        pass
    return out


SALES_TABLE = "dbo.salesforce_bookings_so"


def _sales_where(days, osr, isr, whs, form):
    where = ["order_dt >= DATEADD(day, ?, GETDATE())"]
    params = [-abs(int(days or 90))]
    for col, val in (("os_rep", osr), ("is_rep", isr), ("shp_whs", whs), ("form", form)):
        if val:
            where.append("LTRIM(RTRIM(%s)) = ?" % col)
            params.append(str(val).strip())
    return " AND ".join(where), params


def sales_analytics(days=90, osr=None, isr=None, whs=None, form=None):
    """Visual-dashboard aggregates from salesforce_bookings_so, honoring filters.
    Margin is computed over costed rows only (booked_mtl_cost > 0)."""
    import pyodbc
    cn = pyodbc.connect(CUST_DB_CONN, timeout=60); cur = cn.cursor()
    W, P = _sales_where(days, osr, isr, whs, form)
    names = salesperson_names()

    def rows(sql, extra=()):
        cur.execute(sql, P + list(extra)); return cur.fetchall()

    # KPIs
    k = rows("SELECT SUM(booked_value), SUM(net_wgt), COUNT(DISTINCT so), "
             "SUM(CASE WHEN booked_mtl_cost>0 THEN booked_value END), "
             "SUM(CASE WHEN booked_mtl_cost>0 THEN booked_mtl_cost END) "
             "FROM %s WHERE %s" % (SALES_TABLE, W))[0]
    val, wgt, orders, cval, ccost = (float(k[0] or 0), float(k[1] or 0), int(k[2] or 0),
                                     float(k[3] or 0), float(k[4] or 0))
    kpis = {"value": val, "weight": wgt, "orders": orders,
            "avg_cwt": round(val / (wgt / 100.0), 2) if wgt else None,
            "margin_pct": round(100.0 * (cval - ccost) / cval, 1) if cval else None}

    trend = [{"month": r[0], "val": float(r[1] or 0)} for r in rows(
        "SELECT FORMAT(order_dt,'yyyy-MM') ym, SUM(booked_value) FROM %s WHERE %s "
        "GROUP BY FORMAT(order_dt,'yyyy-MM') ORDER BY ym" % (SALES_TABLE, W))]

    def dim(col, label_names=None, top=None):
        top_sql = ("TOP %d " % top) if top else ""
        rs = rows("SELECT %sLTRIM(RTRIM(%s)) k, SUM(booked_value) v FROM %s WHERE %s "
                  "GROUP BY LTRIM(RTRIM(%s)) ORDER BY SUM(booked_value) DESC"
                  % (top_sql, col, SALES_TABLE, W, col))
        out = []
        for k2, v in rs:
            k2 = (k2 or "(none)")
            lbl = label_names.get(k2, k2) if label_names else k2
            out.append({"key": k2, "label": lbl, "val": float(v or 0)})
        return out

    by_whs = dim("shp_whs")
    by_osr = dim("os_rep", names)
    by_isr = dim("is_rep", names)
    by_form = dim("form")
    top_cust = dim("customer", top=10)

    # heatmap: OSR (rows) x warehouse (cols) booked value
    hm = rows("SELECT LTRIM(RTRIM(os_rep)) r, LTRIM(RTRIM(shp_whs)) c, SUM(booked_value) v "
              "FROM %s WHERE %s GROUP BY LTRIM(RTRIM(os_rep)), LTRIM(RTRIM(shp_whs))" % (SALES_TABLE, W))
    cn.close()
    reps = [d["key"] for d in by_osr][:8]
    cols = [d["key"] for d in by_whs][:6]
    cell = {(r or "(none)", c or "(none)"): float(v or 0) for r, c, v in hm}
    heat = {"rows": [{"key": r, "label": names.get(r, r)} for r in reps], "cols": cols,
            "matrix": [[cell.get((r, c), 0.0) for c in cols] for r in reps]}

    return {"kpis": kpis, "trend": trend, "by_warehouse": by_whs, "by_osr": by_osr,
            "by_isr": by_isr, "by_form": by_form, "top_customers": top_cust, "heatmap": heat,
            "filters": {"days": int(days or 90), "osr": osr, "isr": isr, "whs": whs, "form": form}}


def sales_filter_options():
    """Distinct OSR / ISR / warehouse / form values (last 365d) for the dashboard filters."""
    import pyodbc
    cn = pyodbc.connect(CUST_DB_CONN, timeout=40); cur = cn.cursor()
    names = salesperson_names()
    W = "order_dt >= DATEADD(day, -365, GETDATE())"

    def opts(col, with_names=False):
        cur.execute("SELECT LTRIM(RTRIM(%s)) k, SUM(booked_value) v FROM %s WHERE %s "
                    "AND %s IS NOT NULL AND LTRIM(RTRIM(%s))<>'' "
                    "GROUP BY LTRIM(RTRIM(%s)) ORDER BY SUM(booked_value) DESC"
                    % (col, SALES_TABLE, W, col, col, col))
        return [{"key": r[0], "label": (names.get(r[0], r[0]) if with_names else r[0])}
                for r in cur.fetchall()]
    out = {"osr": opts("os_rep", True), "isr": opts("is_rep", True),
           "whs": opts("shp_whs"), "form": opts("form")}
    cn.close()
    return out


def dashboard(days=90, osr=None, isr=None, whs=None, form=None):
    cfg = get_dash_config()
    conv = _quote_conversion()
    out = {"conversion": conv, "alert_pct": cfg["conv_alert_pct"],
           "alert": (conv["rate"] is not None and conv["rate"] > cfg["conv_alert_pct"]),
           "loss_feedback": _loss_feedback()}
    try:
        out["sales"] = sales_analytics(days, osr, isr, whs, form)
    except Exception as exc:
        out["sales"] = None
        out["sales_error"] = str(exc)[:200]
    try:
        out["options"] = sales_filter_options()
    except Exception as exc:
        out["options"] = None
    try:
        out["quote_history"] = erp_quote_history()
    except Exception as exc:
        out["quote_history"] = None
        out["quote_history_error"] = str(exc)[:200]
    return out



def customers_db_available():
    try:
        import pyodbc
        cn = pyodbc.connect(CUST_DB_CONN, timeout=4); cn.close()
        return True
    except Exception:
        return False


def import_customers_from_db():
    """Pull active customers + credit status live from the ERP (Planning.dbo.transports),
    plus segment from salesforce_billings.dim_seg (matched on customer name).
    Credit code 'H' (hold) -> high risk, else low. Preserves each customer's
    admin-assigned playbook and their earned dynamic margin."""
    import pyodbc
    cn = pyodbc.connect(CUST_DB_CONN, timeout=20)
    cur = cn.cursor()
    cur.execute("""SELECT customer_name, MAX(credit_status)
                   FROM dbo.transports
                   WHERE customer_name IS NOT NULL AND LTRIM(RTRIM(customer_name)) <> ''
                   GROUP BY customer_name""")
    rows = cur.fetchall()
    cn.close()
    resolve_seg = None
    try:
        _, resolve_seg = erp_segment_map()   # resolve(name) -> dim_seg
    except Exception:
        pass
    existing = {c["name"]: c for c in list_customers()}   # also runs the dyn_cwt migration
    conn = connect()
    conn.executescript(SCHEMA)
    n = matched = 0
    for name, cs in rows:
        name = (name or "").strip()
        if not name:
            continue
        credit = "high" if str(cs or "").strip().upper() == "H" else "low"
        ex = existing.get(name)
        seg = resolve_seg(name) if resolve_seg else None
        if seg:
            matched += 1
        else:                                 # keep any prior/admin segment if no ERP match
            seg = ex["segment"] if ex else ""
        # ON CONFLICT preserves each customer's dyn_cwt (their earned margin history).
        conn.execute(
            "INSERT INTO customers (name, segment, playbook, credit) VALUES (?,?,?,?) "
            "ON CONFLICT(name) DO UPDATE SET credit=excluded.credit, segment=excluded.segment",
            (name, seg, ex["playbook"] if ex else 2, credit))
        n += 1
    conn.commit()
    conn.close()
    return {"imported": n, "segmented": matched}


def import_customers_csv(path=None):
    """Bulk-load customers from a CSV with columns name, segment, playbook, credit."""
    import csv
    if path is None:
        path = os.path.join(HERE, "Customers.csv")
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    n = 0
    with open(path, newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            row = {(k or "").strip().lower(): v for k, v in row.items()}
            if not row.get("name"):
                continue
            upsert_customer(row.get("name"), row.get("segment"),
                            row.get("playbook", 2), row.get("credit", "low"))
            n += 1
    return {"imported": n}


if __name__ == "__main__":
    stats = import_from_workbook()
    print("Imported:", stats)
