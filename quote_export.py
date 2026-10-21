"""
Internal Excel export for a priced quote.

Builds a one-row-per-line workbook with product info, weight (lb + tons),
inventory position (on-hand / available / incoming tons and avg cost), the
suggested price, and every pricing adder broken out into its own column
(blank when that adder did not apply).
"""
import io

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import pricing_engine as pe

NAVY = "1F3864"
BLUE = "1D4ED8"
GREY = "F1F5F9"

# Columns: (header, group) -- group drives the light section shading.
COLUMNS = [
    ("Line", "id"),
    ("Thickness", "prod"), ("Form", "prod"), ("Grade", "prod"),
    ("Width (in)", "prod"), ("Length (in)", "prod"), ("Qty", "prod"), ("UOM", "prod"),
    ("Total Weight (lb)", "wt"), ("Total Weight (tons)", "wt"),
    ("On-hand (tons)", "inv"), ("Available (tons)", "inv"),
    ("Avg Cost $/cwt (avail)", "inv"),
    ("Incoming (tons)", "inv"), ("Avg Cost $/cwt (incoming)", "inv"),
    ("Suggested $/cwt", "price"), ("Total Price", "price"),
    ("Base $/cwt", "adder"),
    ("Volume adder", "adder"),
    ("Region - Freight", "adder"), ("Region - Competitive", "adder"),
    ("Inventory position", "adder"),
    ("PB tier", "adder"), ("PB adder", "adder"), ("Segmentation", "adder"),
    ("Credit", "adder"), ("Dynamic margin", "adder"),
    ("Custom-length adder", "adder"), ("Extras", "adder"),
]
HIST_COLUMNS = [
    ("Last sold $/cwt", "hist"), ("Last sold date", "hist"),
    ("Last quote $/cwt", "hist"), ("Last quote outcome", "hist"), ("Last quote date", "hist"),
]
HIST_WIDTHS = [13, 13, 14, 16, 13]
GROUP_FILL = {"prod": "EEF2FF", "wt": "ECFEFF", "inv": "F0FDF4",
              "price": "FEF9C3", "adder": "F8FAFC", "id": "FFFFFF", "hist": "FFF7ED"}


def _adder_column(label):
    """Map a breakdown label to one of the adder columns (or None to skip)."""
    l = (label or "")
    if l.startswith("Base price"):          return "Base $/cwt"
    if l.startswith("Stock adder"):         return "Volume adder"
    if l.startswith("Custom-length"):       return "Custom-length adder"
    if l.startswith("Inventory"):           return "Inventory position"
    if l.startswith("Freight"):             return "Region - Freight"
    if l.startswith("Region (competitive"): return "Region - Competitive"
    if l.startswith("Playbook"):            return "PB adder"
    if l.startswith("Segment"):             return "Segmentation"
    if l.startswith("Credit"):              return "Credit"
    if l.startswith("Dynamic margin"):      return "Dynamic margin"
    if l.startswith("Extra:"):              return "Extras"
    return None


def _form_label(product, inv):
    form = (inv or {}).get("form")
    if form:
        return str(form).upper()                 # resolved CC / CP / FPCO / FP
    pt = (inv or {}).get("product_type")
    if pt:
        return "CTL" if pt.lower() == "ctl" else "Plate"
    th = str((product or {}).get("thickness", "")).strip()
    return "CTL/Sheet" if th.endswith("G") else ("Plate" if th else "")


UOM = {"pieces": "pc", "lbs": "lb", "coil": "coil", "tl": "TL"}


def build_journal_workbook(rows):
    """Excel of the quote journal: one row per generated quote, each stamped with the
    logic version and a snapshot of the config (region rates, adders, credit file)
    that produced it."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quote journal"
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    HDR = ["ID", "Logged At", "Event", "Logic Version", "Customer", "Region", "Lines",
           "Grand Total", "Credit File", "PB1 $", "PB2 $", "PB3 $", "Credit High $",
           "Region Rates (freight/comp)"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HDR))
    ws.cell(1, 1, "Willbanks Pricing App - Quote Journal (logic-stamped history)").font = \
        Font(bold=True, size=13, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 24
    for c, h in enumerate(HDR, 1):
        cell = ws.cell(3, c, h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[3].height = 28
    r = 4
    for row in rows:
        cfg = row.get("config") or {}
        cc = cfg.get("cust_config") or {}
        regs = cfg.get("regions") or []
        reg_str = "; ".join("%s %s/%s" % (x.get("destination"), x.get("freight_cwt"), x.get("comp_cwt"))
                            for x in regs)
        vals = [row.get("id"), row.get("logged_at"), row.get("event"), row.get("logic_version"),
                row.get("customer"), row.get("region"), row.get("lines"), row.get("grand_total"),
                cfg.get("credit_file"), cc.get("pb1_cwt"), cc.get("pb2_cwt"), cc.get("pb3_cwt"),
                cc.get("credit_high_cwt"), reg_str]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if c in (5, 14) else "center", vertical="top")
        ws.cell(r, 8).number_format = '$#,##0.00'
        r += 1
    from openpyxl.utils import get_column_letter
    widths = [6, 18, 9, 13, 26, 14, 6, 13, 24, 7, 7, 7, 11, 52]
    for c, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = wd
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_workbook(payload, result, get_product, playbook=2, pb_adder=0.0, history=None):
    """Return .xlsx bytes for a computed quote.

    payload : the quote request ({lines, customer_name, destination, ...})
    result  : the pricing result from the quote engine ({lines:[...], ...})
    get_product : callable key -> product dict
    playbook / pb_adder : the selected customer's playbook tier (1/2/3) and its $/cwt
                          adder, shown on each priced line.
    history : optional list (aligned with lines) of {sold:[...], quoted:[...]}; when
              given, the last sold/quoted columns are added.
    """
    columns = list(COLUMNS) + (HIST_COLUMNS if history is not None else [])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quote detail"
    headers = [h for h, _ in columns]
    col_of = {h: i + 1 for i, h in enumerate(headers)}
    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # title band
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    cust = (payload.get("customer_name") or "").strip()
    region = (payload.get("destination") or "").strip()
    t = "Willbanks Metals - Internal Quote Detail"
    if cust:
        t += "   |   " + cust
    if region:
        t += "   |   " + region
    c = ws.cell(1, 1, t)
    c.font = Font(bold=True, size=13, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    # header row
    hr = 3
    for i, (h, grp) in enumerate(columns, 1):
        cell = ws.cell(hr, i, h)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[hr].height = 30

    lines_in = payload.get("lines", [])
    lines_out = result.get("lines", [])
    r = hr + 1
    for idx, res in enumerate(lines_out):
        lin = lines_in[idx] if idx < len(lines_in) else {}
        prod = get_product(res.get("key")) or {}
        inv = res.get("inventory") or {}
        row = {h: None for h in headers}
        offsheet = bool(res.get("off_sheet"))
        row["Line"] = idx + 1
        row["Thickness"] = (pe.thickness_label(prod.get("thickness")) if prod.get("thickness")
                            else (lin.get("thickness") or None))
        row["Form"] = ("OFF-SHEET" if offsheet else _form_label(prod, inv))
        row["Grade"] = prod.get("grade") or lin.get("grade")
        row["Width (in)"] = prod.get("width_num") or prod.get("width") or lin.get("width")
        row["Length (in)"] = lin.get("length_in")
        row["Qty"] = res.get("qty") if res.get("qty") is not None else lin.get("qty")
        row["UOM"] = UOM.get((res.get("unit") or lin.get("unit") or "").lower(), res.get("unit"))
        wlb = res.get("weight_lb")
        if wlb is not None:
            row["Total Weight (lb)"] = round(wlb)
            row["Total Weight (tons)"] = round(wlb / 2000.0, 3)

        def tons(x):
            return round(x / 2000.0, 2) if x else None
        row["On-hand (tons)"] = tons(inv.get("sku_ohd_lbs"))
        row["Available (tons)"] = tons(inv.get("sku_available_lbs") or inv.get("grade_available_lbs"))
        row["Avg Cost $/cwt (avail)"] = round(inv["avg_cost_cwt"], 2) if inv.get("avg_cost_cwt") else None
        row["Incoming (tons)"] = tons(inv.get("sku_incoming_lbs") or inv.get("on_order_lbs"))
        # Avg Cost $/cwt (incoming) is not tracked separately in the ERP extract -> left blank

        if res.get("rate_cwt") is not None:
            row["Suggested $/cwt"] = round(res["rate_cwt"], 2)
        if res.get("line_total") is not None:
            row["Total Price"] = round(res["line_total"], 2)

        for b in res.get("breakdown", []):
            col = _adder_column(b.get("label"))
            if col:
                cur = row.get(col) or 0.0
                row[col] = round(cur + float(b.get("cwt") or 0), 2)

        # Show the customer's playbook on every priced line (the PB2 default is $0,
        # so the breakdown omits it -- surface it here so the field is always visible).
        if not res.get("off_sheet") and not res.get("manual_quote") and res.get("rate_cwt") is not None:
            row["PB tier"] = "PB%d" % playbook
            row["PB adder"] = round(float(pb_adder), 2)

        if res.get("manual_quote"):
            row["Suggested $/cwt"] = row["Suggested $/cwt"] or "MANUAL"

        if history is not None:
            h0 = history[idx] if idx < len(history) else {}
            sold = (h0.get("sold") or [None])[0]
            quoted = (h0.get("quoted") or [None])[0]
            if sold:
                row["Last sold $/cwt"] = sold.get("cwt")
                row["Last sold date"] = sold.get("date")
            if quoted:
                row["Last quote $/cwt"] = quoted.get("cwt")
                row["Last quote outcome"] = (quoted.get("outcome") or "").upper()
                row["Last quote date"] = quoted.get("date")

        for h in headers:
            cell = ws.cell(r, col_of[h], row[h])
            cell.border = border
            cell.alignment = Alignment(horizontal="left" if h in ("Form", "Grade", "UOM") else "center")
            if h in ("Avg Cost $/cwt (avail)", "Avg Cost $/cwt (incoming)", "Suggested $/cwt",
                     "Base $/cwt", "Volume adder", "Region - Freight", "Region - Competitive",
                     "Inventory position", "PB adder", "Segmentation", "Credit",
                     "Dynamic margin", "Custom-length adder", "Extras",
                     "Last sold $/cwt", "Last quote $/cwt") and isinstance(row[h], (int, float)):
                cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
            elif h == "Total Price" and isinstance(row[h], (int, float)):
                cell.number_format = '$#,##0.00'
            elif h in ("Total Weight (lb)", "Width (in)", "Length (in)") and isinstance(row[h], (int, float)):
                cell.number_format = '#,##0'
            elif h in ("Total Weight (tons)", "On-hand (tons)", "Available (tons)", "Incoming (tons)") and isinstance(row[h], (int, float)):
                cell.number_format = '#,##0.00'
        r += 1

    # totals row
    gt = result.get("grand_total")
    if gt is not None:
        tc = ws.cell(r, col_of["Total Price"] - 1, "GRAND TOTAL")
        tc.font = Font(bold=True); tc.alignment = Alignment(horizontal="right")
        g = ws.cell(r, col_of["Total Price"], round(gt, 2))
        g.font = Font(bold=True); g.number_format = '$#,##0.00'
        g.fill = PatternFill("solid", fgColor="FEF9C3")

    # group-shade data cells + widths
    for i, (h, grp) in enumerate(columns, 1):
        fill = GROUP_FILL.get(grp)
        for rr in range(hr + 1, r):
            cell = ws.cell(rr, i)
            if cell.value is None or not isinstance(cell.value, str) or True:
                if cell.fill.fgColor.rgb in (None, "00000000"):
                    cell.fill = PatternFill("solid", fgColor=fill)
    widths = [5, 10, 9, 9, 9, 9, 7, 6, 13, 13, 11, 11, 13, 11, 14,
              12, 13, 10, 11, 13, 14, 12, 8, 11, 12, 9, 13, 13, 9] + \
             (HIST_WIDTHS if history is not None else [])
    from openpyxl.utils import get_column_letter
    for i, wd in enumerate(widths[:ncol], 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = ws.cell(hr + 1, 2)
    ws.sheet_view.showGridLines = False

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
